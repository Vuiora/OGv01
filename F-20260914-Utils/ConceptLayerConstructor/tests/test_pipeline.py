import io
import json
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from itertools import pairwise
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from clc.api import create_app
from clc.build import assemble, build_graph, concept_id
from clc.config import Settings
from clc.export import export_graph
from clc.llm import LLM
from clc.models import Citation, Concept, Extraction, Graph, Node, Outline
from clc.parse import chunk_text, parse_document, validate_evidence
from clc.store import Store
from clc.worker import run_one

TEXT = "SCHED_NORMAL 是普通任务的调度策略。CFS 是普通任务公平调度的实现算法。CFS 实现 SCHED_NORMAL 的普通任务调度。"


def settings(tmp_path, **kwargs):
    return Settings(_env_file=None, data_dir=tmp_path, llm_model="test-model", llm_retries=0, **kwargs)


def fixture_reply(payload):
    if "concepts" not in payload:
        return {"concepts": [{"title": title, "kind": kind, "summary": quote, "explanation": quote,
                                  "scope": "Linux 2.6", "evidence": [{"chunk_id": payload["id"], "quote": quote}]}
                              for title, kind, quote in [
                                  ("SCHED_NORMAL", "policy", "SCHED_NORMAL 是普通任务的调度策略。"),
                                  ("CFS", "implementation", "CFS 实现 SCHED_NORMAL 的普通任务调度。")]]}
    ids = {n["title"]: n["id"] for n in payload["concepts"]}
    evidence = next(n for n in payload["concepts"] if n["title"] == "CFS")["evidence"]
    return {"categories": [{"id": "g-scheduling", "title": "普通任务调度", "summary": "策略与实现"}],
            "placements": [{"id": value, "parent_id": "g-scheduling"} for value in ids.values()],
            "relations": [{"source": ids["CFS"], "target": ids["SCHED_NORMAL"], "kind": "implements",
                           "explanation": "CFS 是实现普通任务调度的算法", "evidence": evidence}]}


class FixtureLLM:
    """Test fixture only; production never falls back to fabricated knowledge."""
    def complete(self, schema, task, payload, validate=None):
        result = schema.model_validate(fixture_reply(payload))
        if validate:
            validate(result)
        return result


@pytest.fixture
def prepared(tmp_path):
    folder = tmp_path / "input"
    folder.mkdir()
    (folder / "s001.md").write_text(TEXT, encoding="utf-8")
    return tmp_path, [{"id": "s001", "name": "source.md", "stored_name": "s001.md"}]


def test_grounded_graph_and_all_exports(prepared):
    folder, sources = prepared
    config = settings(folder)
    graph = build_graph("调度策略与 CFS", sources, folder, config, FixtureLLM())
    export_graph(graph, folder / "output", ["md", "docx", "pdf"], config)
    assert len(list((folder / "output/nodes").glob("*.md"))) == len(graph.nodes)
    assert graph.relations[0].kind == "implements"
    assert len({n.kind for n in graph.nodes}) == 3
    assert "CFS" in "\n".join(p.text for p in Document(folder / "output/knowledge.docx").paragraphs)
    pdf = PdfReader(folder / "output/knowledge.pdf")
    assert "CFS" in "".join(p.extract_text() for p in pdf.pages)
    assert "调度" in "".join(p.extract_text() for p in pdf.pages)
    # Every local Markdown link resolves inside the output tree.
    import re
    for path in (folder / "output").rglob("*.md"):
        for link in re.findall(r"\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
            assert (path.parent / link).is_file(), (path, link)


def test_chunk_coverage_and_evidence(tmp_path):
    config = settings(tmp_path, chunk_chars=500, chunk_overlap=40)
    text = "第一段文档。\n\n" * 200
    chunks = chunk_text(text, "s001", config)
    assert chunks[0].start == 0 and chunks[-1].end == len(text)
    assert all(b.start <= a.end for a, b in pairwise(chunks))
    assert all(c.text == text[c.start:c.end] for c in chunks)
    with pytest.raises(ValueError, match="原文引用"):
        validate_evidence([Citation(chunk_id=chunks[0].id, quote="原文没有这个论断")], {c.id: c for c in chunks})


def test_chunk_overlap_does_not_repeat_paragraph_boundary(tmp_path):
    config = settings(tmp_path, chunk_chars=500, chunk_overlap=400)
    text = "a" * 260 + "\n\n" + "b" * 1000
    chunks = chunk_text(text, "s001", config)
    assert len(chunks) <= 10
    assert all(b.end > a.end for a, b in pairwise(chunks))
    assert all(b.start <= a.end for a, b in pairwise(chunks))
    assert chunks[-1].end == len(text)


def test_text_submission_preserves_source_and_uses_configured_limit(tmp_path):
    config = settings(tmp_path, max_source_chars=500_000)
    app = create_app(config)
    text = "    indented code\n" + "x" * 400_001 + "\n\n"
    with TestClient(app) as client:
        result = client.post("/v1/jobs/text", json={"text": text, "title": "  title  ", "formats": ["md"]})
        assert result.status_code == 202
        job = app.state.store.get(result.json()["id"])
        assert job["request"]["title"] == "title"
        source = app.state.store.job_dir(job["id"]) / "input/s001.md"
        assert source.read_text(encoding="utf-8") == text
        assert client.post("/v1/jobs/text", json={"text": "x" * 500_001}).status_code == 413
        assert client.post("/v1/jobs/text", json={"text": "hello", "title": "  "}).status_code == 422


@pytest.mark.parametrize("malformed", [[], {"choices": None}, {"choices": [None]}, {"choices": []}])
def test_llm_retries_malformed_response_envelope(tmp_path, monkeypatch, malformed):
    calls = []

    def handler(request):
        calls.append(request)
        if len(calls) == 1:
            return httpx.Response(200, json=malformed)
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"concepts":[]}'}}]})

    monkeypatch.setattr(time, "sleep", lambda _: None)
    config = settings(tmp_path).model_copy(update={"llm_retries": 1})
    result = LLM(config, httpx.MockTransport(handler)).complete(Extraction, "extract", {})
    assert result.concepts == []
    assert len(calls) == 2


def test_scope_is_part_of_identity():
    concept = Concept.model_validate(fixture_reply({"id": "s001-c0001"})["concepts"][0])
    assert concept_id(concept) != concept_id(concept.model_copy(update={"scope": "Linux 6.12"}))


@pytest.mark.parametrize("parents", [{"a": "b", "b": "a"}, {"a": "missing", "b": None}])
def test_reject_bad_hierarchy(parents):
    with pytest.raises(ValueError):
        Graph(title="test", nodes=[Node(id=k, title=k, kind="category", summary=k,
                                       parent_id=v, synthetic=True) for k, v in parents.items()])


def test_reject_missing_placement():
    node = Node(id="a", title="a", kind="category", summary="a", synthetic=True)
    with pytest.raises(ValueError, match="恰好一次"):
        assemble("t", [node], Outline(placements=[]), {})


def test_llm_http_contract_and_retries(tmp_path, monkeypatch):
    calls = []

    def handler(request):
        calls.append(json.loads(request.content))
        if len(calls) == 1:
            return httpx.Response(429)
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"concepts":[]}'}}]})

    monkeypatch.setattr(time, "sleep", lambda _: None)
    config = settings(tmp_path).model_copy(update={"llm_retries": 1})
    output = LLM(config, httpx.MockTransport(handler)).complete(Extraction, "extract", {"text": "hello"})
    assert output.concepts == [] and len(calls) == 2
    assert calls[-1]["response_format"] == {"type": "json_object"}


def test_api_auth_upload_and_download_guards(tmp_path):
    app = create_app(settings(tmp_path, api_key="secret"))
    with TestClient(app) as client:
        assert client.post("/v1/jobs/text", json={"text": TEXT}).status_code == 401
        headers = {"X-API-Key": "secret"}
        assert client.post("/v1/jobs", headers=headers, files={"files": ("bad.exe", b"x")}).status_code == 415
        assert client.post("/v1/jobs", headers=headers, files={"files": ("bad.pdf", b"not pdf")}).status_code == 422
        assert client.post("/v1/jobs", headers=headers, files={"files": ("bad.docx", b"bad")}).status_code == 422
        assert client.post("/v1/jobs/text", headers=headers, json={"text": "   "}).status_code == 422
        result = client.post("/v1/jobs", headers=headers, files={"files": ("../../danger.md", TEXT.encode())})
        assert result.status_code == 202
        job = result.json()
        stored = app.state.store.get(job["id"])
        assert stored["request"]["sources"][0]["stored_name"] == "s001.md"
        assert client.get(job["bundle_url"], headers=headers).status_code == 409
        assert client.get(f"/v1/jobs/{uuid4()}", headers=headers).status_code == 404


def test_workbench_and_public_configuration(tmp_path):
    app = create_app(settings(tmp_path, api_key="private-service-key", llm_api_key="private-model-key"))
    with TestClient(app) as client:
        page = client.get("/")
        assert page.status_code == 200
        assert "知识层级工作台" in page.text
        assert "text/html" in page.headers["content-type"]
        for asset in ("style.css", "app.js"):
            assert client.get(f"/assets/{asset}").status_code == 200
        health = client.get("/health")
        assert health.json()["auth_required"] is True
        assert health.json()["limits"] == {"files": 10, "upload_mb": 40, "source_chars": 400_000}
        assert "private-service-key" not in health.text
        assert "private-model-key" not in health.text
        assert client.get("/docs").status_code == 200
        assert client.post("/v1/jobs/text", json={"text": TEXT}).status_code == 401


def test_api_upload_limits(tmp_path):
    with TestClient(create_app(settings(tmp_path, max_upload_mb=1))) as client:
        result = client.post("/v1/jobs", files={"files": ("x.md", b"x" * (1024 * 1024 + 1))})
        assert result.status_code == 413
        result = client.post("/v1/jobs", files={"files": ("x.md", b"x" * (3 * 1024 * 1024))})
        assert result.status_code == 413
        assert not list((tmp_path / "jobs").glob("*/input/*"))


def test_local_browser_auto_auth(tmp_path):
    app = create_app(settings(tmp_path, api_key="private-service-key"))
    with TestClient(app, base_url="http://127.0.0.1:8000") as client:
        assert client.post("/v1/jobs/text", json={"text": TEXT}).status_code == 401
        page = client.get("/")
        assert "HttpOnly" in page.headers["set-cookie"]
        assert "SameSite=strict" in page.headers["set-cookie"]
        assert "private-service-key" not in page.text + page.headers["set-cookie"]
        assert client.get("/health").json()["web_auto_auth"] is True
        body = {"text": TEXT, "formats": ["md"]}
        assert client.post("/v1/jobs/text", json=body).status_code == 403
        assert client.post("/v1/jobs/text", json=body, headers={"Origin": "http://evil.example"}).status_code == 403
        response = client.post("/v1/jobs/text", json=body, headers={"Origin": "http://127.0.0.1:8000"})
        assert response.status_code == 202
        assert client.get(response.json()["status_url"]).status_code == 200
    for host, enabled in [("http://remote.example", True), ("http://localhost", False)]:
        app = create_app(settings(tmp_path, api_key="secret", web_auto_auth=enabled))
        with TestClient(app, base_url=host) as client:
            assert "set-cookie" not in client.get("/").headers
            assert client.get("/health").json()["web_auto_auth"] is False
            assert client.post("/v1/jobs/text", json=body).status_code == 401


def test_queue_atomic_claim_and_crash_recovery(tmp_path):
    store = Store(tmp_path)
    job_id = str(uuid4())
    store.enqueue(job_id, {})
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(lambda _: store.claim(30), range(2)))
    assert sum(x is not None for x in results) == 1
    with store.connect() as conn:
        conn.execute("UPDATE jobs SET lease=0 WHERE id=?", (job_id,))
    assert store.claim(30) is None
    assert store.get(job_id)["status"] == "failed"


def test_end_to_end_worker_over_http(tmp_path, monkeypatch):
    # Keep the local fixture server reachable even when Windows has a system proxy.
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            payload = json.loads(body["messages"][1]["content"].split("资料 JSON:\n", 1)[1])
            content = json.dumps(fixture_reply(payload), ensure_ascii=False)
            result = json.dumps({"choices": [{"finish_reason": "stop", "message": {"content": content}}]}).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(result)))
            self.end_headers()
            self.wfile.write(result)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        config = settings(tmp_path, llm_base_url=f"http://127.0.0.1:{server.server_port}/v1")
        app = create_app(config)
        with TestClient(app) as client:
            job = client.post("/v1/jobs/text", json={"title": "调度知识", "text": TEXT}).json()
            assert run_one(app.state.store, config)
            status = client.get(job["status_url"]).json()
            assert status["status"] == "succeeded", status
            assert client.get(job["status_url"] + "/graph").json()["relations"][0]["kind"] == "implements"
            bundle = client.get(job["bundle_url"])
            with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
                assert {"knowledge.pdf", "knowledge.docx", "index.md", "graph.json"} <= set(archive.namelist())
            assert client.get(job["status_url"] + "/files/..%2F..%2Fjobs.sqlite3").status_code == 404
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_n8n_workflow_structure():
    root = Path(__file__).resolve().parents[1]
    for path in (root / "n8n").glob("*.json"):
        flow = json.loads(path.read_text(encoding="utf-8"))
        names = {n["name"] for n in flow["nodes"]}
        assert {"Receive", "Submit", "Status", "Download", "Failed"} <= names
        for outputs in flow["connections"].values():
            for branch in outputs["main"]:
                assert all(edge["node"] in names for edge in branch)
        assert flow["connections"]["Terminal"]["main"][1][0]["node"] == "Wait"


@pytest.mark.skipif(__import__("importlib.util").util.find_spec("docling") is None, reason="optional Docling not installed")
@pytest.mark.parametrize("recover", [True, False])
def test_incomplete_parser_model_recovery(tmp_path, monkeypatch, recover):
    from types import SimpleNamespace

    from docling.datamodel.base_models import ConversionStatus
    from huggingface_hub.errors import IncompleteSnapshotError

    calls = []

    class Converter:
        def __init__(self, **kwargs):
            pass

        def convert(self, *args, **kwargs):
            calls.append(1)
            if not recover or len(calls) == 1:
                raise IncompleteSnapshotError("missing model files", snapshot_path=str(tmp_path))
            return SimpleNamespace(status=ConversionStatus.SUCCESS,
                                   document=SimpleNamespace(export_to_markdown=lambda: TEXT))

    monkeypatch.setattr("docling.document_converter.DocumentConverter", Converter)
    monkeypatch.setattr("clc.parse.time.sleep", lambda _: None)
    if recover:
        assert parse_document(tmp_path / "input.pdf", settings(tmp_path)) == TEXT
        assert len(calls) == 2
    else:
        with pytest.raises(RuntimeError, match="解析模型尚未下载完整"):
            parse_document(tmp_path / "input.pdf", settings(tmp_path))
        assert len(calls) == 3


@pytest.mark.skipif(__import__("importlib.util").util.find_spec("docling") is None, reason="optional Docling not installed")
def test_real_docling_docx(tmp_path):
    path = tmp_path / "input.docx"
    doc = Document()
    doc.add_heading("调度策略", 0)
    doc.add_paragraph(TEXT)
    doc.save(path)
    result = parse_document(path, settings(tmp_path))
    assert "SCHED_NORMAL" in result.replace("\\_", "_") and "CFS" in result
