import json
import os
from pathlib import Path
from xml.sax.saxutils import escape

from markdown_it import MarkdownIt

from clc.config import Settings
from clc.models import Graph

LABELS = {"domain": "领域", "category": "类别", "policy": "策略", "implementation": "实现算法",
          "mechanism": "机制", "concept": "概念", "example": "示例", "is_a": "属于",
          "part_of": "组成部分", "implements": "实现", "handled_by": "由其处理", "depends_on": "依赖",
          "contrasts_with": "对比", "uses": "使用"}


def literal(text: str) -> str:
    return text.replace("\n", " ").replace("\\", "\\\\").replace("[", "\\[").replace("]", "\\]")


def walk(graph: Graph):
    children = {}
    for node in graph.nodes:
        children.setdefault(node.parent_id, []).append(node)

    def visit(parent=None, depth=0):
        for node in children.get(parent, []):
            yield node, depth
            yield from visit(node.id, depth + 1)

    return list(visit())


def node_markdown(node, graph, prefix=""):
    by_id = {n.id: n for n in graph.nodes}
    lines = [f"类型：{LABELS[node.kind]}", f"适用范围：{node.scope}", node.summary]
    if node.synthetic:
        lines.append("此节点是自动生成的阅读目录，不是从原文提取的独立事实。")
    if node.explanation:
        lines.append(node.explanation)
    if node.pitfalls:
        lines.extend(["## 易混点", "\n".join(f"- {p}" for p in node.pitfalls)])
    if node.examples:
        lines.extend(["## 示例", "\n\n".join(node.examples)])
    links = []
    if node.parent_id:
        parent = by_id[node.parent_id]
        links.append(f"- 上级：[{literal(parent.title)}]({prefix}{parent.id}.md)")
    links += [f"- 下级：[{literal(n.title)}]({prefix}{n.id}.md)" for n in graph.nodes if n.parent_id == node.id]
    if links:
        lines.extend(["## 层级导航", "\n".join(links)])
    relations = []
    for rel in graph.relations:
        if node.id in {rel.source, rel.target}:
            source, target = by_id[rel.source], by_id[rel.target]
            relations.append(f"- [{literal(source.title)}]({prefix}{source.id}.md) → {LABELS[rel.kind]} → "
                             f"[{literal(target.title)}]({prefix}{target.id}.md)：{rel.explanation}")
            relations.extend(f"  - 关系依据 {c.chunk_id}：{literal(c.quote)}" for c in rel.evidence)
    if relations:
        lines.extend(["## 相关关系", "\n".join(relations)])
    if node.evidence:
        lines.append("## 原文依据")
        for cite in node.evidence:
            source_id = cite.chunk_id.split("-c")[0]
            source_prefix = "../sources/" if not prefix else "sources/"
            lines.append(f"- 来源 [{source_id}]({source_prefix}{source_id}.md)，分块 `{cite.chunk_id}`\n\n"
                         + "\n".join("> " + line for line in cite.quote.splitlines()))
    return "\n\n".join(lines) + "\n"


def blocks(markdown):
    tokens = MarkdownIt().parse(markdown)
    heading, listed = 0, 0
    for token in tokens:
        if token.type == "heading_open":
            heading = int(token.tag[1:])
        elif token.type == "heading_close":
            heading = 0
        elif token.type == "list_item_open":
            listed += 1
        elif token.type == "list_item_close":
            listed -= 1
        elif token.type == "inline":
            yield ("heading" if heading else "list" if listed else "paragraph", heading, token.children or [])
        elif token.type in {"fence", "code_block"}:
            yield ("code", 0, token.content)


def text_content(children):
    return "".join("\n" if t.type in {"softbreak", "hardbreak"} else t.content
                   for t in children if t.type in {"text", "code_inline", "softbreak", "hardbreak", "image"})


def export_docx(markdown: str, path: Path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ["Normal", "Title", "List Bullet", *[f"Heading {i}" for i in range(1, 10)]]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    for kind, level, value in blocks(markdown):
        if kind == "code":
            for line in value.splitlines():
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                run = para.add_run(line)
                run.font.name, run.font.size = "Consolas", Pt(9)
            continue
        if kind == "heading":
            para = doc.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 9)}")
        else:
            para = doc.add_paragraph(style="List Bullet" if kind == "list" else "Normal")
        bold = italic = False
        for token in value:
            if token.type == "strong_open":
                bold = True
            elif token.type == "strong_close":
                bold = False
            elif token.type == "em_open":
                italic = True
            elif token.type == "em_close":
                italic = False
            elif token.type in {"text", "code_inline", "image", "softbreak", "hardbreak"}:
                run = para.add_run("\n" if token.type in {"softbreak", "hardbreak"} else token.content)
                run.bold, run.italic = bold, italic
                if token.type == "code_inline":
                    run.font.name = "Consolas"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(path)


def export_pdf(markdown: str, path: Path, settings: Settings):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    candidates = [settings.pdf_font] if settings.pdf_font else [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc",
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
    ]
    font_path = next((p for p in candidates if p and p.is_file()), None)
    if font_path is None:
        raise RuntimeError("PDF 导出需要中文 TrueType 字体，请设置 CLC_PDF_FONT；Linux 可安装 fonts-wqy-zenhei")
    font = "CLCFont"
    pdfmetrics.registerFont(TTFont(font, str(font_path)))
    base = {"fontName": font, "textColor": colors.black, "wordWrap": "CJK", "alignment": TA_LEFT,
            "splitLongWords": True, "allowWidows": 0, "allowOrphans": 0}
    styles = {"paragraph": ParagraphStyle("body", fontSize=10.5, leading=17, spaceAfter=8, **base),
              "list": ParagraphStyle("list", fontSize=10.5, leading=17, leftIndent=12, spaceAfter=5, **base),
              "code": ParagraphStyle("code", fontSize=9, leading=14, leftIndent=12, spaceAfter=8, **base)}
    story = []
    for kind, level, value in blocks(markdown):
        if kind == "heading":
            size = {1: 24, 2: 18, 3: 14}.get(level, 12)
            style = ParagraphStyle(f"h{level}", fontSize=size, leading=size * 1.4,
                                   spaceBefore=16, spaceAfter=8, keepWithNext=True, **base)
        else:
            style = styles[kind]
        plain = value if kind == "code" else text_content(value)
        content = escape(plain).replace("\n", "<br/>")
        if kind == "list":
            content = "• " + content
        story.append(Paragraph(content or " ", style))

    def page_number(canvas, doc):
        canvas.setFont(font, 9)
        canvas.drawRightString(A4[0] - 48, 28, str(doc.page))

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=48, rightMargin=48,
                      topMargin=44, bottomMargin=44).build(story, onFirstPage=page_number, onLaterPages=page_number)


def export_graph(graph: Graph, output: Path, formats: list[str], settings: Settings):
    nodes_dir = output / "nodes"
    nodes_dir.mkdir(parents=True, exist_ok=True)
    intro = ("本文依据输入资料组织知识点，用目录呈现阅读层级，用独立关系说明实现、依赖和对比。"
             "每个知识点均保留原文依据；自动分类和解释仍需人工复核，适用范围以来源为准。")
    index = [f"# {literal(graph.title)}", intro, "## 知识目录"]
    book = [f"# {literal(graph.title)}", intro]
    for node, depth in walk(graph):
        metadata = {"id": node.id, "kind": node.kind, "parent_id": node.parent_id, "synthetic": node.synthetic}
        header = "---\n" + "\n".join(f"{k}: {json.dumps(v, ensure_ascii=False)}" for k, v in metadata.items()) + "\n---\n\n"
        (nodes_dir / f"{node.id}.md").write_text(header + f"# {literal(node.title)}\n\n" + node_markdown(node, graph), encoding="utf-8")
        index.append("  " * depth + f"- [{literal(node.title)}](nodes/{node.id}.md) · {LABELS[node.kind]}")
        body = node_markdown(node, graph, "nodes/")
        # Keep aggregate section levels valid while the precise full hierarchy remains in index + graph.
        heading_level = min(depth + 2, 5)
        body = body.replace("\n## ", "\n" + "#" * min(heading_level + 1, 6) + " ")
        book.extend(["#" * heading_level + " " + literal(node.title), body])
    (output / "index.md").write_text("\n\n".join(index[:3]) + "\n\n" + "\n".join(index[3:]) + "\n", encoding="utf-8")
    markdown = "\n\n".join(book)
    (output / "knowledge.md").write_text(markdown, encoding="utf-8")
    if "docx" in formats:
        export_docx(markdown, output / "knowledge.docx")
    if "pdf" in formats:
        export_pdf(markdown, output / "knowledge.pdf", settings)
